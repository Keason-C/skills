"""把 Word / Excel / PDF 写到磁盘上的小工具。

只有两个用户:测试夹具(`tests/conftest.py`)和演示语料生成器
(`scripts/make_demo_corpus.py`)。两边造的是不同的语料,但"怎么写出一个
真的 .docx"是同一件事,不该抄两遍。

只在开发期使用 —— 运行时的系统只**读**文档,从不写。
"""

from __future__ import annotations

import io
from pathlib import Path


def write_docx(path: Path, paragraphs: list[str]) -> None:
    from docx import Document

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))


def write_xlsx(path: Path, rows: list[list[str]], sheet_title: str = "Sheet1") -> None:
    from openpyxl import Workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.active.title = sheet_title
    for row in rows:
        wb.active.append(row)
    wb.save(str(path))


def write_pdf(path: Path, lines: list[str], *, font_size: int = 12) -> None:
    """带文字层的 PDF。中文用 CID 字体,pypdf 抽得出来。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    path.parent.mkdir(parents=True, exist_ok=True)
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.setFont("STSong-Light", font_size)
    y = 780
    for line in lines:
        c.drawString(56, y, line)
        y -= font_size + 6
    c.save()
    path.write_bytes(buf.getvalue())


def write_scanned_pdf(path: Path) -> None:
    """**没有文字层**的 PDF —— 模拟扫描/拍照的合同。"""
    from reportlab.pdfgen import canvas

    path.parent.mkdir(parents=True, exist_ok=True)
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.rect(56, 520, 480, 260, stroke=1, fill=0)
    c.line(80, 700, 500, 700)
    c.line(80, 660, 500, 660)
    c.save()
    path.write_bytes(buf.getvalue())


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
